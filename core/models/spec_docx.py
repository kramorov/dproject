# core/models/spec_docx.py
"""Рендер спецификации каталога в .docx через docxtpl.

Единый шаблон для всех моделей каталога: контекст строится методом
``TemplateMixin.get_spec_doc_context()`` (секции характеристик текстом,
дефолтное изображение, техдокументация и сертификаты со ссылками на
скачивание — полный вариант и «ужатый»), а этот модуль раскладывает его
по .docx-шаблону.

Использование:

    from core.models.spec_docx import render_spec_docx
    render_spec_docx(item, '/tmp/spec.docx', base_url='https://example.com')
"""

import io
import os
import re
from html import escape as _html_escape

from django.conf import settings
from docx.shared import Mm

try:
    from docxtpl import DocxTemplate, InlineImage
except Exception:  # docxtpl не установлен — рендер недоступен, но импорт не падает
    DocxTemplate = InlineImage = None


DEFAULT_TEMPLATE_PATH = os.path.join(
    settings.BASE_DIR, 'core', 'templates', 'docx', 'specification_template.docx',
)


# Хром документа («Характеристики», «Артикул» и т.п.) по локалям.
# Ключи соответствуют плейсхолдерам {{ chrome.* }} в build_spec_template().
SPEC_CHROME = {
    'ru': {
        'article': 'Артикул',
        'characteristics': 'Характеристики',
        'tech_docs': 'Техническая документация',
        'certs': 'Сертификаты',
    },
    'en': {
        'article': 'Article',
        'characteristics': 'Specifications',
        'tech_docs': 'Technical documentation',
        'certs': 'Certificates',
    },
    'cn': {
        'article': '型号',
        'characteristics': '技术参数',
        'tech_docs': '技术文档',
        'certs': '证书',
    },
}


def _fetch_image_bytes(image: dict):
    """Байты изображения: локальный media-файл → URL (S3 presigned и т.п.)."""
    url = image.get('url')
    if not url:
        return None

    # Локальное хранилище: /media/<path> → MEDIA_ROOT/<path>
    if str(url).startswith('/media/'):
        local = os.path.join(settings.MEDIA_ROOT, str(url)[len('/media/'):])
        if os.path.exists(local):
            with open(local, 'rb') as fh:
                return _normalize_image(fh.read())

    try:
        import urllib.request
        with urllib.request.urlopen(str(url), timeout=30) as resp:
            return _normalize_image(resp.read())
    except Exception:
        return None


def _normalize_image(data):
    """Привести изображение к PNG/JPEG (python-docx не умеет WebP и т.п.)."""
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(data))
    except Exception:
        return data
    if (img.format or '').upper() in ('PNG', 'JPEG'):
        return data
    out = io.BytesIO()
    img.save(out, format='PNG')
    return out.getvalue()


def _to_floating_anchor(pic_xml: str) -> str:
    """Превратить ``<wp:inline>`` в плавающий ``<wp:anchor>``.

    Картинка прижимается к правому верхнему углу поля с обтеканием текста
    (wrapSquare, bothSides). Сохраняет xmlns-декларации исходного inline-тега.
    """
    xml = re.sub(r'<wp:inline\b', '<wp:anchor', pic_xml, count=1)
    xml = xml.replace('</wp:inline>', '</wp:anchor>')
    xml = re.sub(
        r'<wp:anchor\b([^>]*)>',
        r'<wp:anchor\1 distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        r'relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"'
        r'><wp:simplePos x="0" y="0"/>'
        r'<wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>'
        r'<wp:positionV relativeFrom="margin"><wp:align>top</wp:align></wp:positionV>',
        xml,
        count=1,
    )
    xml = re.sub(
        r'(<wp:extent\b[^>]*/>)',
        r'\1<wp:wrapSquare wrapText="bothSides"/>',
        xml,
        count=1,
    )
    return xml


if InlineImage is not None:
    class FloatingImage(InlineImage):
        """InlineImage, рендерящийся как плавающая картинка с обтеканием."""

        def _insert_image(self):
            pic = self.tpl.current_rendering_part.new_pic_inline(
                self.image_descriptor, self.width, self.height,
            ).xml
            pic = _to_floating_anchor(pic)
            return (
                "</w:t></w:r><w:r><w:drawing>%s</w:drawing></w:r><w:r>"
                '<w:t xml:space="preserve">' % pic
            )
else:
    FloatingImage = None


def _link_run(rid: str, label: str):
    """XML одного гиперлинка внутри абзаца."""
    return (
        '<w:hyperlink r:id="%s" w:tgtFrame="_blank">'
        '<w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
        '<w:t xml:space="preserve">%s</w:t></w:r></w:hyperlink>'
    ) % (rid, _html_escape(label))


def _build_rich_links(entries, tpl):
    """К имени добавляет две ссылки (полный и «ужатый»), как сырой run-XML.

    Плейсхолдер ``{{ d.rich }}`` лежит внутри ``<w:t>…</w:t>``, поэтому строка
    закрывает текущий run, вставляет свои runs/hyperlinks и переоткрывает run —
    ровно как docxtpl.InlineImage.
    """
    for entry in entries:
        name = entry.get('name') or ''
        parts = ['</w:t></w:r>']
        parts.append(
            '<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r>'
            % _html_escape(name)
        )
        if entry.get('url_full'):
            parts.append(_link_run(tpl.build_url_id(entry['url_full']), ' Скачать'))
        if entry.get('url_compressed'):
            parts.append(_link_run(tpl.build_url_id(entry['url_compressed']), ' Скачать (сжат)'))
        parts.append('<w:r><w:t xml:space="preserve">')
        entry['rich'] = ''.join(parts)
    return entries


def build_spec_template(path: str):
    """Создать .docx-шаблон спецификации с Jinja-плейсхолдерами docxtpl.

    Структура: заголовок/артикул → изображение → характеристики (по группам)
    → техдокументация → сертификаты. Один шаблон на все модели каталога.
    Пустые секции скрываются через ``{%p if %}``.

    Блочные теги используют форму ``{%p ... %}`` (paragraph-loop): абзац с тегом
    вырезается docxtpl целиком, поэтому в документе не остаётся пустых строк.

    ВНИМАНИЕ: плейсхолдеры ``{{ image }}`` и ``{{ d.rich }}``/``{{ c.rich }}``
    должны оставаться в одном run (как генерит этот метод). Не редактируйте
    шаблон вручную — при изменении структуры перегенерируйте его через
    ``build_spec_template``.
    """
    from docx import Document

    doc = Document()

    doc.add_heading('{{ item.title }}', level=1)
    doc.add_paragraph('{{ chrome.article }}: {{ item.code }}')

    # Изображение — плейсхолдер в отдельном абзаце
    doc.add_paragraph('{{ image }}')

    # ── Характеристики ──
    doc.add_paragraph('{%p if spec_groups %}')
    doc.add_heading('{{ chrome.characteristics }}', level=2)
    doc.add_paragraph('{%p for g in spec_groups %}')
    doc.add_heading('{{ g.title }}', level=3)
    doc.add_paragraph('{%p for row in g.rows %}')
    doc.add_paragraph('{{ row.label }}: {{ row.text }}')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endif %}')

    # ── Техническая документация ──
    doc.add_paragraph('{%p if tech_docs %}')
    doc.add_heading('{{ chrome.tech_docs }}', level=2)
    doc.add_paragraph('{%p for d in tech_docs %}')
    doc.add_paragraph('{{ d.rich }}')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endif %}')

    # ── Сертификаты ──
    doc.add_paragraph('{%p if certs %}')
    doc.add_heading('{{ chrome.certs }}', level=2)
    doc.add_paragraph('{%p for c in certs %}')
    doc.add_paragraph('{{ c.rich }}')
    doc.add_paragraph('{%p endfor %}')
    doc.add_paragraph('{%p endif %}')

    d = os.path.dirname(path)
    if d:
        os.makedirs(d, exist_ok=True)
    doc.save(path)
    return path


def render_spec_docx_bytes(item, template_path: str = None, base_url: str = None,
                          image_width_mm: float = 90.0, image_float: bool = True,
                          locale=None) -> bytes:
    """Рендер спецификации ``item`` в байты .docx (для HTTP-скачивания).

    ``image_float`` — True: картинка плавает в правом верхнем углу с обтеканием
    текста; False: обычная inline-картинка в потоке.
    ``locale`` — локаль данных и хрома документа (ru/en/cn).
    """
    from ..utils.localization import DEFAULT_LOCALE

    if DocxTemplate is None:
        raise RuntimeError('docxtpl не установлен — установите docxtpl')

    locale = locale or DEFAULT_LOCALE
    ctx = item.get_spec_doc_context(base_url=base_url, locale=locale)
    ctx['chrome'] = SPEC_CHROME.get(locale, SPEC_CHROME[DEFAULT_LOCALE])

    tpl_path = template_path or DEFAULT_TEMPLATE_PATH
    if not os.path.exists(tpl_path):
        build_spec_template(tpl_path)

    tpl = DocxTemplate(tpl_path)

    # ── Изображение ──
    image = ctx.get('image')
    if isinstance(image, dict):
        data = _fetch_image_bytes(image)
        if data:
            img_cls = FloatingImage if image_float else InlineImage
            ctx['image'] = img_cls(tpl, io.BytesIO(data), width=Mm(image_width_mm))
        else:
            ctx['image'] = ''
    else:
        ctx['image'] = ''

    # ── Ссылки на скачивание ──
    ctx['tech_docs'] = _build_rich_links(ctx.get('tech_docs', []), tpl)
    ctx['certs'] = _build_rich_links(ctx.get('certs', []), tpl)

    tpl.render(ctx)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def render_spec_docx(item, output_path: str, template_path: str = None,
                     base_url: str = None, image_width_mm: float = 90.0,
                     image_float: bool = True, locale=None):
    """Рендер спецификации ``item`` в файл .docx.

    ``template_path`` — путь к .docx-шаблону; если не задан и файла нет —
    шаблон создаётся автоматически (``build_spec_template``).
    """
    data = render_spec_docx_bytes(item, template_path, base_url, image_width_mm,
                                  image_float, locale=locale)
    with open(output_path, 'wb') as fh:
        fh.write(data)
    return output_path
