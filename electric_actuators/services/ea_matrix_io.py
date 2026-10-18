# electric_actuators/services/ea_matrix_io.py
"""
Экспорт/импорт матрицы сигналов серии ЭП в Excel (xlsx).

Экспорт — **полный каркас рабочей таблицы**: все модели серии × все пары
(БУ × напряжение × режим), встречающиеся в матрице серии/бренда или в
разрешённых БУ×напряжение серии (`AllowedControlUnitOption`), × все опции
каждого активного слота (каталог `EaSlotOption`). Ячейки без строки —
пустая `availability` (импорт их пропускает): таблицу можно дозаполнять в Excel
и заливать обратно (upsert на уровне корпуса).

Строка файла = опция: series/item/body, БУ, напряжение, режим, слот, вид
(signal/heater/none), код опции (`slot:code`), код сигнала/формы, питание
обогрева, availability (std/opt/na/пусто), is_default, bundle, устройство,
encoding/ext_encoding. Идентификация — по `option_code` (fallback — signal_code).
"""
from django.db.models import Q

from electric_actuators.models import (
    EaSignalMatrixEntry, EaSlotOption, ElectricActuatorModelLine,
    ElectricActuatorBody, ElectricActuatorModelLineItem,
)
from electric_actuators.services.ea_signal_resolver import resolve_matrix
from params.models import ControlUnitInstalledOption, PowerSupplies

FIELDS = ['series_code', 'item_code', 'body_code', 'control_unit',
          'power_supply', 'mode', 'slot', 'kind', 'option_code',
          'signal_code', 'form_code', 'heater_supply', 'availability',
          'is_default', 'bundle', 'device', 'encoding', 'ext_encoding']


def _scope_body_code(body):
    if not body:
        return ''
    return body.code or body.name or ''


def _option_key(entry):
    """Идентификатор опции строки матрицы: каталог-код или fallback."""
    if entry.option_id:
        return entry.option.code
    if entry.feedback_signal_id:
        return f'{entry.slot.code}:{entry.feedback_signal.code}'
    if entry.control_signal_id:
        return f'{entry.slot.code}:{entry.control_signal.code}'
    if entry.protocol_id:
        return f'{entry.slot.code}:{entry.protocol.code}'
    if entry.heater_supply_id:
        return f'{entry.slot.code}:{entry.heater_supply.code}'
    return f'{entry.slot.code}:{entry.name or ""}'


def _work_row(series_code, item_code, body_code, cu_code, ps_code, mode,
              slot_code, kind, option_code, signal_code, form_code,
              heater_code, availability, is_default, bundle, device_code,
              encoding, ext_encoding):
    return {
        'series_code': series_code,
        'item_code': item_code,
        'body_code': body_code,
        'control_unit': cu_code,
        'power_supply': ps_code,
        'mode': mode,
        'slot': slot_code,
        'kind': kind,
        'option_code': option_code,
        'signal_code': signal_code,
        'form_code': form_code,
        'heater_supply': heater_code,
        'availability': availability,
        'is_default': bool(is_default),
        'bundle': bundle or '',
        'device': device_code,
        'encoding': encoding or '',
        'ext_encoding': ext_encoding or '',
    }


def _entry_codes(entry):
    """Коды сигнала/формы/обогрева/устройства строки матрицы."""
    signal_code = form_code = heater_code = device_code = ''
    if entry.feedback_signal_id:
        signal_code = entry.feedback_signal.code
        form_code = entry.feedback_signal.form.code if entry.feedback_signal.form_id else ''
    elif entry.control_signal_id:
        signal_code = entry.control_signal.code
        form_code = entry.control_signal.form.code if entry.control_signal.form_id else ''
    elif entry.protocol_id:
        signal_code = entry.protocol.code
    if entry.heater_supply_id:
        heater_code = entry.heater_supply.code
    if entry.device_id:
        device_code = entry.device.code
    return signal_code, form_code, heater_code, device_code


def _series_matrix_qs(line):
    """Все строки матрицы, относящиеся к серии (любой скоуп)."""
    return EaSignalMatrixEntry.objects.filter(
        Q(brand=line.brand) | Q(model_line=line)
        | Q(body__model_line=line) | Q(model_line_item__model_line=line))


def _series_pairs(line):
    """Пары (cu_code, ps_code, mode) серии: матрица + разрешённые БУ×напряжение
    (AllowedControlUnitOption — новая связка; режим для БУ — both)."""
    pairs = set()
    qs = _series_matrix_qs(line)
    for r in qs.values('control_unit__code', 'power_supply__code',
                       'mode').distinct():
        pairs.add((r['control_unit__code'], r['power_supply__code'], r['mode']))
    from electric_actuators.models import AllowedControlUnitOption
    for a in AllowedControlUnitOption.objects.filter(
            model_line=line, is_active=True,
            power_supply__isnull=False).select_related(
            'control_unit', 'power_supply'):
        pairs.add((a.control_unit.code, a.power_supply.code, 'both'))
    return pairs


def export_series_rows(model_line_code):
    """Каркас матрицы серии: список dict для xlsx (пустые ячейки — '')."""
    line = ElectricActuatorModelLine.objects.get(code=model_line_code)
    items = list(ElectricActuatorModelLineItem.objects.filter(
        model_line=line).select_related('body').order_by('code'))
    options = {}
    for opt in EaSlotOption.objects.filter(
            is_active=True, slot__is_active=True).select_related(
            'slot', 'feedback_signal', 'control_signal', 'protocol',
            'heater_supply', 'device'):
        options.setdefault(opt.slot.code, []).append(opt)
    cus = {c.code: c for c in ControlUnitInstalledOption.objects.all()}
    pss = {p.code: p for p in PowerSupplies.objects.all()}
    matrix_qs = _series_matrix_qs(line)
    pairs = sorted(_series_pairs(line))

    rows = []
    for item in items:
        body_code = _scope_body_code(item.body)
        item_code = item.code or item.name
        for cu_code, ps_code, mode in pairs:
            cu, ps = cus.get(cu_code), pss.get(ps_code)
            if not cu or not ps:
                continue
            # Эффективные строки (резолвер) и запреты (na)
            res = resolve_matrix(item, cu, ps, mode)
            effective = {}
            for slot in res['slots']:
                for srow in slot['rows']:
                    entry = EaSignalMatrixEntry.objects.get(pk=srow['id'])
                    effective[_option_key(entry)] = dict(
                        availability=srow['availability'],
                        is_default=srow['is_default'],
                        bundle=srow['bundle'] or '',
                        entry=entry)
            na_keys = set()
            for entry in matrix_qs.filter(availability='na').select_related(
                    'slot', 'control_unit', 'power_supply'):
                if (entry.control_unit.code != cu_code
                        or entry.power_supply.code != ps_code):
                    continue
                if entry.mode not in (mode, 'both'):
                    continue
                if entry.body_id and entry.body_id != (item.body_id or -1):
                    continue
                if (entry.model_line_item_id
                        and entry.model_line_item_id != item.id):
                    continue
                na_keys.add(_option_key(entry))

            emitted = set()
            for slot_code, opts in options.items():
                for opt in opts:
                    key = opt.code
                    data = effective.get(key)
                    if key in na_keys:
                        availability = 'na'
                        is_default = False
                        bundle = ''
                        entry = None
                    elif data:
                        availability = data['availability']
                        is_default = data['is_default']
                        bundle = data['bundle']
                        entry = data['entry']
                    else:
                        availability = ''
                        is_default = False
                        bundle = ''
                        entry = None
                    if opt.kind == 'pair':
                        pair = opt.code.split(':', 1)[1]  # 'END_SPDT'
                        kind, contact = pair.split('_', 1)
                        sig_codes = {'OPEN': f'FB_{kind}_OPEN_{contact}',
                                     'CLOSE': f'FB_{kind}_CLOSE_{contact}'}
                        for pos, sig_code in sig_codes.items():
                            emitted.add(key)
                            rows.append(_work_row(
                                line.code, item_code, body_code, cu_code,
                                ps_code, mode, slot_code, 'signal', key,
                                sig_code, '', '', availability, is_default,
                                bundle, '', '', ''))
                    elif opt.kind == 'signal':
                        sig = opt.feedback_signal or opt.control_signal or opt.protocol
                        form_code = ''
                        if sig and getattr(sig, 'form_id', None):
                            form_code = sig.form.code
                        emitted.add(key)
                        rows.append(_work_row(
                            line.code, item_code, body_code, cu_code,
                            ps_code, mode, slot_code, 'signal', key,
                            sig.code if sig else '', form_code, '',
                            availability, is_default, bundle,
                            opt.device.code if opt.device_id else '',
                            entry.encoding if entry else '',
                            entry.ext_encoding if entry else ''))
                    elif opt.kind == 'heater':
                        emitted.add(key)
                        rows.append(_work_row(
                            line.code, item_code, body_code, cu_code,
                            ps_code, mode, slot_code, 'heater', key, '',
                            '', opt.heater_supply.code, availability,
                            is_default, bundle, '',
                            entry.encoding if entry else '',
                            entry.ext_encoding if entry else ''))
                    else:  # none
                        emitted.add(key)
                        rows.append(_work_row(
                            line.code, item_code, body_code, cu_code,
                            ps_code, mode, slot_code, 'none', key, '', '',
                            '', availability, is_default, bundle,
                            opt.device.code if opt.device_id else '',
                            entry.encoding if entry else '',
                            entry.ext_encoding if entry else ''))
            # Строки вне каталога (не потерять данные)
            for key, data in effective.items():
                if key in emitted:
                    continue
                entry = data['entry']
                sig_code, form_code, heater_code, device_code = _entry_codes(
                    entry)
                kind = ('heater' if entry.heater_supply_id
                        else 'none' if not (entry.feedback_signal_id
                                            or entry.control_signal_id
                                            or entry.protocol_id)
                        else 'signal')
                rows.append(_work_row(
                    line.code, item_code, body_code, cu_code, ps_code, mode,
                    entry.slot.code, kind, key, sig_code, form_code,
                    heater_code, data['availability'], data['is_default'],
                    data['bundle'], device_code, entry.encoding,
                    entry.ext_encoding))
    return rows


def write_xlsx(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ea_signal_matrix'
    ws.append(FIELDS)
    for row in rows:
        ws.append([row.get(f, '') for f in FIELDS])
    wb.save(path)
    return len(rows)


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    for i, values in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        rec = dict(zip(FIELDS, values))
        if not rec.get('slot'):
            continue
        if not rec.get('availability') or rec['availability'] == '':
            continue  # пустая ячейка каркаса — пропускаем
        rows.append(rec)
    return rows


def apply_import(rows):
    """Upsert строк матрицы из файла (скоуп — корпус по body_code)."""
    from params.models import ActuatorHeaterSupply
    from electric_actuators.models import (EaSignalSlot, EaFeedbackSignal,
                                           EaControlSignal, EaCommProtocol,
                                           EaFeedbackDevice, EaSlotOption)
    stats = {'created': 0, 'updated': 0, 'errors': []}
    slots = {s.code: s for s in EaSignalSlot.objects.all()}
    opts = {o.code: o for o in EaSlotOption.objects.all()}
    for rec in rows:
        try:
            if not rec.get('availability') or rec['availability'] == '':
                continue  # пустая ячейка каркаса — пропускаем
            body = None
            if rec.get('body_code'):
                body = ElectricActuatorBody.objects.filter(
                    code=rec['body_code']).first() or \
                    ElectricActuatorBody.objects.filter(
                        name=rec['body_code']).first()
            item = ElectricActuatorModelLineItem.objects.filter(
                code=rec.get('item_code')).first()
            if body is None and item is not None:
                body = item.body
            from params.models import ControlUnitInstalledOption, PowerSupplies
            cu = ControlUnitInstalledOption.objects.get(code=rec['control_unit'])
            ps = PowerSupplies.objects.get(code=rec['power_supply'])
            slot = slots.get(rec['slot'])
            if slot is None:
                raise ValueError(f'Слот {rec["slot"]} не найден')
            avail = rec['availability'] or 'opt'
            kw = dict(mode=rec['mode'], availability=avail,
                      is_default=(rec.get('is_default') in (True, 'True', 1)
                                  or avail == 'std'),
                      encoding=rec.get('encoding') or '',
                      ext_encoding=rec.get('ext_encoding') or '')
            option = opts.get(rec.get('option_code') or '')
            bundle = rec.get('bundle') or ''
            lookup = dict(control_unit=cu, power_supply=ps, mode=rec['mode'],
                          slot=slot)
            if body is not None:
                lookup['body'] = body
            elif item is not None:
                lookup['model_line_item'] = item
            else:
                raise ValueError('Нет body_code/item_code для строки')
            kind = rec.get('kind') or ('heater' if rec.get('heater_supply')
                                       else 'none' if not rec.get('signal_code')
                                       else 'signal')
            if option is not None:
                kw['option'] = option
            if kind == 'heater':
                hs = ActuatorHeaterSupply.objects.get(
                    code=rec['heater_supply'])
                lookup['heater_supply'] = hs
                obj, created = EaSignalMatrixEntry.objects.update_or_create(
                    defaults=kw, **lookup)
            elif kind == 'none':
                kw['bundle'] = bundle or None
                kw['name'] = (option.name if option and option.name
                              else rec.get('option_code', ''))
                device = EaFeedbackDevice.objects.filter(
                    code=rec.get('device')).first() if rec.get('device') else None
                kw['device'] = device
                # Имя/опция — часть идентичности none-строки (иначе разные
                # none-опции одного слота затирают друг друга)
                lookup['name'] = kw['name']
                obj, created = EaSignalMatrixEntry.objects.update_or_create(
                    defaults=kw, **lookup)
            else:
                sig_code = rec.get('signal_code')
                fb = EaFeedbackSignal.objects.filter(code=sig_code).first()
                ctrl = EaControlSignal.objects.filter(code=sig_code).first()
                proto = EaCommProtocol.objects.filter(code=sig_code).first()
                if fb:
                    lookup['feedback_signal'] = fb
                elif ctrl:
                    lookup['control_signal'] = ctrl
                elif proto:
                    lookup['protocol'] = proto
                else:
                    raise ValueError(f'Сигнал {sig_code} не найден')
                device = EaFeedbackDevice.objects.filter(
                    code=rec.get('device')).first() if rec.get('device') else None
                kw['device'] = device
                if not bundle and option and option.kind == 'pair':
                    pair = option.code.split(':', 1)[1]
                    scope_code = (body.code or body.name) if body else (
                        item.code or item.name)
                    kw['bundle'] = f'{pair}:{scope_code}:{rec["mode"]}'
                else:
                    kw['bundle'] = bundle or None
                obj, created = EaSignalMatrixEntry.objects.update_or_create(
                    defaults=kw, **lookup)
            stats['created' if created else 'updated'] += 1
        except Exception as e:  # noqa: BLE001 — отчёт по строкам
            stats['errors'].append(f'{rec.get("item_code")}/{rec.get("slot")}: {e}')
    return stats
