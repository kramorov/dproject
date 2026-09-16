# commercial/management/commands/create_quotation_template.py
"""
Создать дефолтный .docx-шаблон КП (шапка).

Использование:
    python manage.py create_quotation_template

Шаблон — обычный Word-файл с Jinja2-тегами (docxtpl). Содержит шапку документа
(заголовок, номер, дату, клиента). Табличная часть и спецификации добавляются
программно (см. commercial/services/quotation.py).

Файл можно открыть в Word и редактировать: добавить логотип, реквизиты,
дополнительный текст, изменить вёрстку шапки.
"""
from pathlib import Path

from django.core.management.base import BaseCommand
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


class Command(BaseCommand):
    help = 'Создать дефолтный шаблон КП (quotation_template.docx).'

    def handle(self, *args, **options):
        doc = Document()

        normal = doc.styles['Normal']
        normal.font.name = 'Arial'
        normal.font.size = Pt(11)

        # ── Шапка ──
        title = doc.add_heading('Коммерческое предложение', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('№ {{ number }} от {{ date }}').bold = True

        doc.add_paragraph('Клиент: {{ customer }}')
        doc.add_paragraph('Корзина: {{ cart_name }}')
        doc.add_paragraph()

        # Далее (таблица + спецификации) добавляются программно.

        out = Path(__file__).resolve().parent.parent.parent / 'templates' / 'quotation_template.docx'
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))

        self.stdout.write(self.style.SUCCESS(f'Template created: {out}'))
