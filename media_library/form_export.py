"""
Экспорт структуры распознавания (JSON-дерева) в Word/Excel.

Входной формат — дерево узлов:
    {
      "items": [
        {"kind": "section", "title": "...", "children": [ ... ]},
        {"kind": "text", "text": "..."},
        {"kind": "table", "variant": "fv",
         "rows": [ {"field": "...", "value": "..."}, ... ]},
        {"kind": "table", "variant": "grid",
         "headers": ["a", "b"], "rows": [["x", "y"], ...]}
      ]
    }

Поддерживается и устаревший узел {"kind": "field", "field": ..., "values": [...]}.
"""

from io import BytesIO
from typing import Any , Dict , List


def _flatten_value(value : Dict[str , Any]) -> str :
    """Привести значение (устаревший формат field/values) к строке."""
    kind = value.get('type' , 'text')
    if kind == 'checkboxes' :
        selected = value.get('selected') or []
        if selected :
            return ', '.join(str(s) for s in selected)
        options = value.get('options') or []
        checked = [o.get('label' , '') for o in options if o.get('checked')]
        return ', '.join(str(c) for c in checked if c)
    return (value.get('text') or '').strip()


def _iter_rows(items : List[Dict[str , Any]] ,
               section_path : str = '' ,
               out : list | None = None) -> list :
    """Развернуть дерево в плоский список для листа «Структура».

    Элементы списка:
        ('pair',  section_path, field, value)          — строка колонок
        ('grid',  section_path, headers, rows)         — классическая таблица
    """
    if out is None :
        out = []
    for item in items :
        kind = item.get('kind')
        if kind == 'section' :
            title = item.get('title') or ''
            path = (section_path + ' / ' + title) if section_path else title
            _iter_rows(item.get('children') or [] , path , out)
        elif kind == 'text' :
            out.append(('pair' , section_path , '' , (item.get('text') or '').strip()))
        elif kind == 'field' :  # устаревший формат
            field = item.get('field') or ''
            values = item.get('values') or []
            if not values :
                out.append(('pair' , section_path , field , ''))
            for v in values :
                out.append(('pair' , section_path , field , _flatten_value(v)))
        elif kind == 'table' :
            variant = item.get('variant') or 'fv'
            if variant == 'grid' :
                out.append((
                    'grid' , section_path ,
                    [str(h or '') for h in (item.get('headers') or [])] ,
                    [[c for c in row] for row in (item.get('rows') or [])] ,
                ))
            else :
                for row in item.get('rows') or [] :
                    out.append((
                        'pair' , section_path ,
                        row.get('field' , '') or '' , row.get('value' , '') or '' ,
                    ))
    return out


def _sheet_title(number : int) -> str :
    return f'Таблица {number}'


def structure_to_xlsx(structure : Dict[str , Any]) -> bytes :
    """Структура → Excel (bytes).

    Лист «Структура»: колонки Раздел / Поле / Значение.
    Классические таблицы выносятся на отдельные листы «Таблица N»
    (первая строка — заголовки, если заданы, иначе «Столбец N»).
    """
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Структура'
    ws.append(['Раздел' , 'Поле' , 'Значение'])

    grid_no = 0
    for entry in _iter_rows(structure.get('items') or []) :
        if entry[0] == 'grid' :
            _ , path , headers , rows = entry
            grid_no += 1
            sheet = wb.create_sheet(_sheet_title(grid_no))
            width = max([len(headers)] + [len(r) for r in rows] , default=0)
            if headers :
                sheet.append(headers + ['' for _ in range(max(0 , width - len(headers)))])
            elif width :
                sheet.append([f'Столбец {i + 1}' for i in range(width)])
            for r in rows :
                sheet.append([c for c in r])
            ws.append([path , '' , f'{_sheet_title(grid_no)} → лист «{_sheet_title(grid_no)}»'])
        else :
            _ , path , field , value = entry
            ws.append([path , field , value])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fill_row(cells , values , ncols : int) -> None :
    """Записать строку значений в ячейки таблицы python-docx."""
    for i in range(ncols) :
        cells[i].text = str(values[i]) if i < len(values) and values[i] is not None else ''


def _render_docx_items(doc , items : List[Dict[str , Any]] , level : int = 0) -> None :
    for item in items :
        kind = item.get('kind')
        if kind == 'section' :
            doc.add_heading(item.get('title') or '' , level=min(level + 1 , 4))
            _render_docx_items(doc , item.get('children') or [] , level + 1)
        elif kind == 'text' :
            text = (item.get('text') or '').strip()
            if text :
                doc.add_paragraph(text)
        elif kind == 'field' :  # устаревший формат
            field = item.get('field') or ''
            values = item.get('values') or []
            if values :
                for v in values :
                    flat = _flatten_value(v)
                    doc.add_paragraph(f'{field}: {flat}' if field else flat)
            else :
                doc.add_paragraph(field)
        elif kind == 'table' :
            variant = item.get('variant') or 'fv'
            rows = item.get('rows') or []
            if not rows :
                continue
            if variant == 'grid' :
                headers = [str(h or '') for h in (item.get('headers') or [])]
                ncols = max([len(headers)] + [len(r) for r in rows] , default=0)
                if ncols <= 0 :
                    continue
                table = doc.add_table(rows=1 , cols=ncols)
                table.style = 'Table Grid'
                if headers :
                    _fill_row(table.rows[0].cells , headers , ncols)
                    for r in rows :
                        _fill_row(table.add_row().cells , r , ncols)
                else :
                    _fill_row(table.rows[0].cells , rows[0] , ncols)
                    for r in rows[1:] :
                        _fill_row(table.add_row().cells , r , ncols)
            else :
                table = doc.add_table(rows=1 , cols=2)
                table.style = 'Table Grid'
                table.rows[0].cells[0].text = 'Поле'
                table.rows[0].cells[1].text = 'Значение'
                for row in rows :
                    cells = table.add_row().cells
                    cells[0].text = row.get('field' , '') or ''
                    cells[1].text = row.get('value' , '') or ''


def structure_to_docx(structure : Dict[str , Any]) -> bytes :
    """Структура → Word (bytes): заголовки разделов, абзацы текста, таблицы."""
    from docx import Document
    doc = Document()
    _render_docx_items(doc , structure.get('items') or [])
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
