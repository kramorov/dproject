# media_library/table_extractor.py
"""
Распознавание таблиц на изображении через img2table — инструмент медиабиблиотеки.

На входе: изображение (bytes или объект с атрибутом ``media_file``, например
``MediaLibraryItem``). На выходе: список распознанных таблиц (pandas.DataFrame)
и готовый Excel-файл (bytes).

OCR-движок подключаемый: имя бэкенда → класс ``img2table.ocr.*``. Импорты
img2table / pandas / numpy — ленивые, чтобы отсутствие библиотеки не роняло
импорт приложения Django. Выбор конкретного OCR делается на этапе вызова;
бэкенд по умолчанию задаётся настройкой ``MEDIA_TABLE_OCR_BACKEND``.
"""

import logging
from io import BytesIO
from typing import Any , Dict , List , Optional

logger = logging.getLogger(__name__)

# Имя бэкенда → (модуль, класс, дефолтные kwargs).
# Список расширяется по мере подключения новых OCR-движков.
#
# Требования бэкендов (img2table 2.x):
#   tesseract — внешний бинарь tesseract.exe (UB-Mannheim), без pip-пакетов
#   rapidocr   — pip install rapidocr (onnxruntime, без внешних бинарей)
#   surya      — pip install surya-ocr (torch, тяжёлый)
#   paddle     — pip install paddleocr paddlepaddle
#   easyocr    — pip install easyocr (torch, тяжёлый)
#   doctr      — pip install python-doctr
OCR_BACKENDS = {
    'tesseract' : ('img2table.ocr' , 'TesseractOCR' , {'lang' : 'rus+eng' , 'n_threads' : 1}) ,
    'rapidocr' : ('img2table.ocr' , 'RapidOCR' , {}) ,
    'surya' : ('img2table.ocr' , 'SuryaOCR' , {}) ,
    'paddle' : ('img2table.ocr' , 'PaddleOCR' , {'lang' : 'en'}) ,
    'easyocr' : ('img2table.ocr' , 'EasyOCR' , {'lang' : ['ru' , 'en']}) ,
    'doctr' : ('img2table.ocr' , 'DocTR' , {}) ,
}


class TableExtractionError(RuntimeError) :
    """Ошибка распознавания таблицы."""


def available_backends() -> List[str] :
    """Список зарегистрированных OCR-бэкендов."""
    return sorted(OCR_BACKENDS.keys())


def default_backend() -> str :
    """Бэкенд по умолчанию: настройка MEDIA_TABLE_OCR_BACKEND или 'tesseract'."""
    try :
        from django.conf import settings
        backend = getattr(settings , 'MEDIA_TABLE_OCR_BACKEND' , None)
        if backend and backend in OCR_BACKENDS :
            return backend
    except Exception :
        pass
    return 'tesseract'


def _import_img2table_image() :
    """Ленивый импорт img2table.document.Image."""
    try :
        from img2table.document import Image as Img2TableImage
        return Img2TableImage
    except ImportError as e :
        raise TableExtractionError(
            'img2table не установлен. Установите: pip install img2table'
        ) from e


def get_ocr_instance(backend : str , **kwargs) :
    """Создать OCR-объект img2table по имени бэкенда.

    kwargs переопределяют дефолтные параметры бэкенда (например ``lang``).
    """
    if backend not in OCR_BACKENDS :
        raise TableExtractionError(
            f'Неизвестный OCR-бэкенд "{backend}". Доступные: {available_backends()}'
        )

    module_name , class_name , defaults = OCR_BACKENDS[backend]
    params = dict(defaults)
    params.update({k : v for k , v in kwargs.items() if v is not None})

    import importlib
    try :
        module = importlib.import_module(module_name)
        ocr_cls = getattr(module , class_name)
        return ocr_cls(**params)
    except ImportError as e :
        raise TableExtractionError(
            f'OCR-бэкенд "{backend}" требует img2table и зависимости OCR: {e}'
        ) from e
    except Exception as e :
        raise TableExtractionError(
            f'Не удалось инициализировать OCR "{backend}": {e}'
        ) from e


def _read_image_bytes(media_file) -> bytes :
    """Прочитать байты из FieldFile (работает для local и Cloud.ru)."""
    if media_file is None :
        raise TableExtractionError('У медиаэлемента нет файла (media_file).')
    try :
        media_file.open('rb')
        try :
            data = media_file.read()
        finally :
            media_file.close()
    except Exception as e :
        raise TableExtractionError(f'Не удалось прочитать файл: {e}') from e

    if not data :
        raise TableExtractionError('Файл пуст.')
    return data


def crop_image_bytes(data : bytes ,
                     x : int , y : int , w : int , h : int) -> bytes :
    """Обрезать изображение по прямоугольной области (пиксели исходника).

    Координаты клампится по границам изображения; результат отдаётся как PNG.
    """
    from PIL import Image

    try :
        img = Image.open(BytesIO(data))
        x1 = max(0 , int(x))
        y1 = max(0 , int(y))
        x2 = min(img.width , x1 + int(w))
        y2 = min(img.height , y1 + int(h))
        if x2 <= x1 or y2 <= y1 :
            raise TableExtractionError('Пустая область выделения.')
        cropped = img.crop((x1 , y1 , x2 , y2))
        buf = BytesIO()
        cropped.save(buf , 'PNG')
        return buf.getvalue()
    except TableExtractionError :
        raise
    except Exception as e :
        raise TableExtractionError(f'Не удалось обрезать изображение: {e}') from e


def extract_tables(item_or_bytes ,
                   ocr_backend : Optional[str] = None ,
                   ocr_kwargs : Optional[Dict[str , Any]] = None ,
                   borderless_tables : bool = False ,
                   implicit_rows : bool = True ,
                   min_confidence : int = 50 ,
                   max_workers : int = 1) -> List :
    """
    Распознать таблицы на изображении.

    Args:
        item_or_bytes: MediaLibraryItem (или любой объект с .media_file) либо bytes.
        ocr_backend: имя OCR-бэкенда из OCR_BACKENDS (None → default_backend()).
        ocr_kwargs: параметры, передаваемые в конструктор OCR.
        borderless_tables: искать таблицы без видимых линий.
        implicit_rows: достраивать неявные строки.
        min_confidence: минимальная уверенность детекции таблицы (0-100).
        max_workers: число потоков OCR.

    Returns:
        Список объектов ExtractedTable (у каждого есть .df, .bbox, .html_repr()).
    """
    Img2TableImage = _import_img2table_image()

    if hasattr(item_or_bytes , 'media_file') :
        data = _read_image_bytes(item_or_bytes.media_file)
    elif isinstance(item_or_bytes , (bytes , bytearray)) :
        data = bytes(item_or_bytes)
    else :
        raise TableExtractionError(
            'Ожидается MediaLibraryItem или bytes с изображением.'
        )

    if not data :
        raise TableExtractionError('Пустые данные изображения.')

    ocr_backend = ocr_backend or default_backend()
    ocr = get_ocr_instance(ocr_backend , **(ocr_kwargs or {}))

    try :
        document = Img2TableImage(src=data)
        tables = document.extract_tables(
            ocr=ocr ,
            implicit_rows=implicit_rows ,
            borderless_tables=borderless_tables ,
            min_confidence=min_confidence ,
            max_workers=max_workers ,
        )
    except Exception as e :
        raise TableExtractionError(f'Ошибка распознавания: {e}') from e

    logger.info(
        f'Распознано таблиц: {len(tables)} (ocr={ocr_backend})'
    )
    return tables or []


def _to_json_safe(value) :
    """Привести значение ячейки к JSON-совместимому типу."""
    import numpy as np
    import pandas as pd

    if value is None :
        return None
    if isinstance(value , pd.Timestamp) :
        return value.isoformat()
    if isinstance(value , (np.integer ,)) :
        return int(value)
    if isinstance(value , (np.floating ,)) :
        f = float(value)
        return None if np.isnan(f) else f
    if isinstance(value , (np.bool_ ,)) :
        return bool(value)
    if isinstance(value , float) and np.isnan(value) :
        return None
    if isinstance(value , (int , float , str , bool)) :
        return value
    return str(value)


def _bbox_to_list(bbox) -> Optional[List[int]] :
    """BBox (dataclass) → [x1, y1, x2, y2] (Python int, не numpy).

    BBox не итерируемый, а координаты бывают numpy-целыми — без int()
    json.dump падает с «int64 is not JSON serializable».
    """
    if bbox is None :
        return None
    return [int(bbox.x1) , int(bbox.y1) , int(bbox.x2) , int(bbox.y2)]


def table_to_dict(table , index : int = 0 ,
                  use_first_row_as_header : bool = True) -> Dict[str , Any] :
    """Сериализовать ExtractedTable в dict (для JSON-превью).

    ``use_first_row_as_header`` продвигает первую строку в ``columns``
    (img2table 2.x возвращает шапку первой строкой ``df``).
    """
    df = table.df
    rows = [
        [_to_json_safe(v) for v in row]
        for row in df.values.tolist()
    ]
    columns = [str(c) for c in df.columns.tolist()]
    if use_first_row_as_header and rows :
        columns = [str(v) if v is not None else '' for v in rows[0]]
        rows = rows[1:]

    return {
        'index' : index ,
        'n_columns' : len(columns) ,
        'n_rows' : len(rows) ,
        'columns' : columns ,
        'rows' : rows ,
        'bbox' : _bbox_to_list(getattr(table , 'bbox' , None)) ,
        'title' : getattr(table , 'title' , None) ,
    }


def tables_to_preview(tables : List ,
                      use_first_row_as_header : bool = True) -> List[Dict[str , Any]] :
    """Список таблиц в виде dict для JSON-ответа."""
    return [
        table_to_dict(t , i , use_first_row_as_header=use_first_row_as_header)
        for i , t in enumerate(tables)
    ]


def tables_to_excel_bytes(tables : List ,
                          sheet_name_prefix : str = 'Table') -> bytes :
    """Собрать распознанные таблицы в один Excel-файл (bytes).

    Одна таблица → лист ``Table``, несколько → ``Table 1``, ``Table 2``, ...
    """
    if not tables :
        raise TableExtractionError('Таблицы не найдены.')

    import pandas as pd

    buf = BytesIO()
    with pd.ExcelWriter(buf , engine='openpyxl') as writer :
        if len(tables) == 1 :
            tables[0].df.to_excel(writer , index=False , sheet_name=sheet_name_prefix)
        else :
            for i , table in enumerate(tables , start=1) :
                table.df.to_excel(
                    writer , index=False ,
                    sheet_name=f'{sheet_name_prefix} {i}' ,
                )
    return buf.getvalue()


def _records_to_text(records) -> str :
    """Собрать полный текст из OCR-записей img2table (по строкам)."""
    from collections import defaultdict

    groups = defaultdict(list)
    for page_records in records.values() :
        for record in page_records :
            groups[record.get('parent')].append(record)

    lines = []
    for group in groups.values() :
        try :
            y1 = min(float(w.get('y1' , 0)) for w in group)
            x1 = min(float(w.get('x1' , 0)) for w in group)
        except (TypeError , ValueError) :
            y1 = x1 = 0.0
        value = ' '.join(
            str(w.get('value' , '') or '')
            for w in sorted(group , key=lambda w : float(w.get('x1' , 0) or 0))
        ).strip()
        if value :
            lines.append((y1 , x1 , value))

    lines.sort(key=lambda t : (t[0] , t[1]))
    return '\n'.join(v for _ , _ , v in lines)


def _recognize_page(data : bytes ,
                    ocr ,
                    Img2TableImage ,
                    borderless_tables : bool ,
                    implicit_rows : bool ,
                    min_confidence : int ,
                    max_workers : int ,
                    use_first_row_as_header : bool ,
                    include_text : bool) -> Dict[str , Any] :
    """Распознать одну страницу (один проход OCR): текст + таблицы."""
    if not data :
        raise TableExtractionError('Пустые данные изображения.')

    document = Img2TableImage(src=data)

    ocr_data = None
    text = ''
    if include_text :
        try :
            ocr_data = ocr.of(document)
            if ocr_data is not None :
                text = _records_to_text(ocr_data.records)
        except Exception as e :
            logger.warning('Не удалось извлечь текст: %s' , e)
            ocr_data = None

    try :
        if ocr_data is not None :
            document = Img2TableImage(src=data , ocr_data=ocr_data)
            tables = document.extract_tables(
                ocr=None ,
                implicit_rows=implicit_rows ,
                borderless_tables=borderless_tables ,
                min_confidence=min_confidence ,
                max_workers=max_workers ,
            )
        else :
            tables = Img2TableImage(src=data).extract_tables(
                ocr=ocr ,
                implicit_rows=implicit_rows ,
                borderless_tables=borderless_tables ,
                min_confidence=min_confidence ,
                max_workers=max_workers ,
            )
    except Exception as e :
        raise TableExtractionError(f'Ошибка распознавания: {e}') from e

    tables = tables or []
    return {
        'text' : text ,
        'tables' : tables_to_preview(
            tables , use_first_row_as_header=use_first_row_as_header
        ) ,
    }


def recognize_pages(page_images : List[bytes] ,
                    ocr_backend : Optional[str] = None ,
                    ocr_kwargs : Optional[Dict[str , Any]] = None ,
                    borderless_tables : bool = False ,
                    implicit_rows : bool = True ,
                    min_confidence : int = 50 ,
                    max_workers : int = 1 ,
                    include_text : bool = True ,
                    use_first_row_as_header : bool = True) -> Dict[str , Any] :
    """Распознать одну или несколько страниц (OCR-движок создаётся один раз).

    Returns:
        {'pages': [{'n', 'text', 'tables'}, ...], 'page_count', 'tables', 'text'}
    """
    Img2TableImage = _import_img2table_image()
    ocr_backend = ocr_backend or default_backend()
    ocr = get_ocr_instance(ocr_backend , **(ocr_kwargs or {}))

    pages = []
    for i , data in enumerate(page_images , start=1) :
        page = _recognize_page(
            data=data , ocr=ocr , Img2TableImage=Img2TableImage ,
            borderless_tables=borderless_tables ,
            implicit_rows=implicit_rows ,
            min_confidence=min_confidence ,
            max_workers=max_workers ,
            use_first_row_as_header=use_first_row_as_header ,
            include_text=include_text ,
        )
        page['n'] = i
        pages.append(page)

    total_tables = sum(len(p.get('tables') or []) for p in pages)
    logger.info(
        'Распознано страниц: %d, таблиц: %d (ocr=%s)' ,
        len(pages) , total_tables , ocr_backend ,
    )
    return {
        'pages' : pages ,
        'page_count' : len(pages) ,
        'tables' : [t for p in pages for t in p.get('tables') or []] ,
        'text' : '\n\n'.join(
            p.get('text') for p in pages if (p.get('text') or '').strip()
        ).strip() ,
    }


def result_to_excel_bytes(result : Dict[str , Any] ,
                          include_text : bool = True) -> bytes :
    """Сериализованный результат (pages) → Excel-файл (bytes)."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    pages = result.get('pages') or []
    has_any = False

    for page in pages :
        pn = page.get('n' , 0)
        if include_text and (page.get('text') or '').strip() :
            ws = wb.create_sheet(f'Page {pn} - Text')
            for line in page['text'].strip().split('\n') :
                ws.append([line])
            has_any = True
        for t in (page.get('tables') or []) :
            ws = wb.create_sheet(f'Page {pn} - Table {t.get("index" , 0) + 1}')
            ws.append([str(c) for c in (t.get('columns') or [])])
            for row in (t.get('rows') or []) :
                ws.append([None if c is None else c for c in row])
            has_any = True

    if not has_any :
        ws = wb.create_sheet('Empty')
        ws.append(['Ничего не распознано'])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_to_docx_bytes(result : Dict[str , Any]) -> bytes :
    """Сериализованный результат (pages) → Word-документ (.docx, bytes)."""
    from docx import Document

    doc = Document()
    pages = result.get('pages') or []
    multi = len(pages) > 1

    for page in pages :
        pn = page.get('n' , 0)
        if multi :
            doc.add_heading(f'Страница {pn}' , level=1)

        text = (page.get('text') or '').strip()
        if text :
            for para in text.split('\n') :
                if para.strip() :
                    doc.add_paragraph(para.strip())

        for t in (page.get('tables') or []) :
            doc.add_heading(f'Таблица {t.get("index" , 0) + 1}' , level=2)
            cols = [str(c) for c in (t.get('columns') or [])]
            rows = t.get('rows') or []
            ncols = len(cols) or (len(rows[0]) if rows else 0)
            if ncols == 0 :
                continue

            table = doc.add_table(rows=1 , cols=ncols)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for j in range(ncols) :
                hdr[j].text = cols[j] if j < len(cols) else ''

            for row in rows :
                cells = table.add_row().cells
                for j in range(ncols) :
                    val = row[j] if j < len(row) else ''
                    cells[j].text = '' if val is None else str(val)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
